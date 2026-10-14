"""Construcción de reportes y exportación a Excel y PDF.

Todos los reportes se describen con la misma estructura:
    {titulo, subtitulo, empresa, secciones: [{titulo, columnas, filas, totales}]}
columnas: [{clave, titulo, tipo}] con tipo en texto | numero | moneda | porcentaje
"""
import io
from datetime import datetime

import bd
import calculos
import configuracion as cfg

TIPOS = {
    "hoja_costos": "Hoja de costos",
    "materiales": "Reporte de materiales",
    "mano_obra": "Reporte de mano de obra",
    "cif": "Reporte de CIF",
    "produccion": "Reporte de producción",
    "variaciones": "Reporte de variaciones",
    "estado_costos": "Estado de costos de producción",
}


def col(clave, titulo, tipo="texto"):
    return {"clave": clave, "titulo": titulo, "tipo": tipo}


def _sumar(filas, claves):
    return {k: round(sum(float(f.get(k) or 0) for f in filas), 2) for k in claves}


def _encabezado(eid, periodo_id=None):
    e = bd.uno("SELECT * FROM empresas WHERE id=%s", (eid,))
    p = bd.uno("SELECT * FROM periodos WHERE id=%s", (periodo_id,)) if periodo_id else None
    return e, p


def _ordenes(eid, periodo_id, producto_id=None):
    sql = "SELECT id FROM ordenes_produccion WHERE empresa_id=%s AND periodo_id=%s"
    params = [eid, periodo_id]
    if producto_id:
        sql += " AND producto_id=%s"
        params.append(producto_id)
    return [calculos.calcular_orden(f["id"]) for f in bd.consultar(sql + " ORDER BY numero", params)]


def construir(tipo, eid, periodo_id=None, producto_id=None, orden_id=None):
    if tipo not in TIPOS:
        raise ValueError("Tipo de reporte desconocido")
    e, p = _encabezado(eid, periodo_id)
    rep = {"tipo": tipo, "titulo": TIPOS[tipo], "empresa": e["razon_social"] if e else "",
           "nit": e["nit"] if e else "", "periodo": p["nombre"] if p else "",
           "generado": datetime.now().strftime("%Y-%m-%d %H:%M"), "secciones": []}
    S = rep["secciones"]

    if tipo == "hoja_costos":
        h = calculos.calcular_orden(orden_id)
        o = h["orden"]
        rep["periodo"] = o["periodo"]
        rep["subtitulo"] = "Orden %s · %s %s · %s" % (o["numero"], o["producto_codigo"], o["producto"], o["fecha"])
        S.append({"titulo": "Datos de la orden", "columnas": [col("c", "Concepto"), col("v", "Valor", "numero")],
                  "filas": [{"c": "Unidades terminadas", "v": o["cantidad_terminada"]},
                            {"c": "Unidades en proceso", "v": o["cantidad_en_proceso"]},
                            {"c": "% avance materiales", "v": o["avance_materiales"]},
                            {"c": "% avance conversión", "v": o["avance_conversion"]},
                            {"c": "Unidades equivalentes materiales", "v": h["unidades_equiv_mp"]},
                            {"c": "Unidades equivalentes conversión", "v": h["unidades_equiv_conv"]}]})
        filas = list(h["resumen"])
        S.append({"titulo": "Costo estándar vs. costo real", "columnas": [
            col("elemento", "Elemento"), col("estandar_unitario", "Estándar unitario", "moneda"),
            col("estandar_total", "Estándar total", "moneda"), col("real_unitario", "Real unitario", "moneda"),
            col("real_total", "Real total", "moneda"), col("diferencia", "Diferencia", "moneda")],
            "filas": filas,
            "totales": {"elemento": "TOTAL", "estandar_unitario": h["total_estandar_unitario"],
                        "estandar_total": h["total_estandar"], "real_unitario": h["total_real_unitario"],
                        "real_total": h["total_real"], "diferencia": h["diferencia_total"]}})
        S.append({"titulo": "Materiales", "columnas": [
            col("materia_prima", "Materia prima"), col("cantidad_estandar", "Cant. estándar", "numero"),
            col("precio_estandar", "Precio est.", "moneda"), col("cantidad_real", "Cant. real", "numero"),
            col("precio_real", "Precio real", "moneda"), col("variacion_precio", "Var. precio", "moneda"),
            col("variacion_cantidad", "Var. cantidad", "moneda")],
            "filas": h["materiales"], "totales": {"materia_prima": "TOTAL", **_sumar(h["materiales"], ["variacion_precio", "variacion_cantidad"])}})
        S.append({"titulo": "Mano de obra directa", "columnas": [
            col("cargo", "Cargo"), col("horas_estandar", "Horas estándar", "numero"),
            col("tarifa_estandar", "Tarifa est.", "moneda"), col("horas_reales", "Horas reales", "numero"),
            col("tarifa_real", "Tarifa real", "moneda"), col("variacion_tarifa", "Var. tarifa", "moneda"),
            col("variacion_eficiencia", "Var. eficiencia", "moneda")],
            "filas": h["mano_obra"], "totales": {"cargo": "TOTAL", **_sumar(h["mano_obra"], ["variacion_tarifa", "variacion_eficiencia"])}})
        c = h["cif"]
        S.append({"titulo": "Costos indirectos de fabricación", "columnas": [col("c", "Concepto"), col("v", "Valor", "numero")],
                  "filas": [{"c": "Base de aplicación", "v": c["base_aplicacion"]},
                            {"c": "Tasa predeterminada total", "v": c["tasa_total"]},
                            {"c": "Base estándar permitida", "v": c["base_estandar_permitida"]},
                            {"c": "Base real", "v": c["base_real"]},
                            {"c": "CIF aplicado", "v": c["cif_aplicado"]},
                            {"c": "CIF real prorrateado", "v": c["cif_real_prorrateado"]},
                            {"c": "Diferencia", "v": c["diferencia"]}]})
        return rep

    if not periodo_id:
        raise ValueError("Seleccione un periodo")
    rep["subtitulo"] = "Periodo %s (%s a %s)" % (p["nombre"], p["fecha_inicio"], p["fecha_fin"])

    if tipo == "materiales":
        filas = []
        for h in _ordenes(eid, periodo_id, producto_id):
            for m in h["materiales"]:
                filas.append({"orden": h["orden"]["numero"], **m})
        S.append({"titulo": "Consumo de materiales por orden", "columnas": [
            col("orden", "Orden"), col("materia_prima", "Materia prima"),
            col("cantidad_estandar", "Cant. estándar", "numero"), col("cantidad_real", "Cant. real", "numero"),
            col("costo_estandar", "Costo estándar", "moneda"), col("costo_real", "Costo real", "moneda"),
            col("variacion_precio", "Var. precio", "moneda"), col("variacion_cantidad", "Var. cantidad", "moneda")],
            "filas": filas, "totales": {"orden": "TOTAL", **_sumar(filas, ["costo_estandar", "costo_real", "variacion_precio", "variacion_cantidad"])}})
        catalogo = bd.consultar(
            "SELECT m.codigo, m.descripcion, u.codigo AS unidad, m.precio_estandar, pr.nombre AS proveedor, cc.nombre AS centro"
            " FROM materias_primas m JOIN unidades_medida u ON u.id=m.unidad_id LEFT JOIN proveedores pr ON pr.id=m.proveedor_id"
            " LEFT JOIN centros_costo cc ON cc.id=m.centro_costo_id WHERE m.empresa_id=%s ORDER BY m.codigo", (eid,))
        S.append({"titulo": "Catálogo de materias primas", "columnas": [
            col("codigo", "Código"), col("descripcion", "Descripción"), col("unidad", "Unidad"),
            col("precio_estandar", "Precio estándar", "moneda"), col("proveedor", "Proveedor"), col("centro", "Centro de costo")],
            "filas": catalogo})

    elif tipo == "mano_obra":
        filas = []
        for h in _ordenes(eid, periodo_id, producto_id):
            for m in h["mano_obra"]:
                filas.append({"orden": h["orden"]["numero"], **m})
        S.append({"titulo": "Mano de obra directa por orden", "columnas": [
            col("orden", "Orden"), col("cargo", "Cargo"), col("horas_estandar", "Horas est.", "numero"),
            col("horas_reales", "Horas reales", "numero"), col("costo_estandar", "Costo estándar", "moneda"),
            col("costo_real", "Costo real", "moneda"), col("variacion_tarifa", "Var. tarifa", "moneda"),
            col("variacion_eficiencia", "Var. eficiencia", "moneda")],
            "filas": filas, "totales": {"orden": "TOTAL", **_sumar(filas, ["horas_estandar", "horas_reales", "costo_estandar", "costo_real", "variacion_tarifa", "variacion_eficiencia"])}})
        f = calculos.factores_prestacionales(eid)
        trab = bd.consultar("SELECT t.*, c.nombre AS cargo FROM trabajadores t JOIN cargos c ON c.id=t.cargo_id"
                            " WHERE t.empresa_id=%s ORDER BY t.nombre", (eid,))
        for t in trab:
            t.update(calculos.costo_trabajador(t, f))
        S.append({"titulo": "Nómina y tarifa real por hora", "columnas": [
            col("documento", "Documento"), col("nombre", "Trabajador"), col("cargo", "Cargo"),
            col("salario_mensual", "Salario", "moneda"), col("factor_prestacional", "Factor prest. %", "numero"),
            col("costo_mensual_total", "Costo mensual", "moneda"), col("horas_mes", "Horas/mes", "numero"),
            col("tarifa_real_hora", "Tarifa real/hora", "moneda")],
            "filas": trab, "totales": {"documento": "TOTAL", **_sumar(trab, ["salario_mensual", "costo_mensual_total"])}})

    elif tipo == "cif":
        filas = bd.consultar("SELECT c.codigo, c.concepto, c.comportamiento, cc.nombre AS centro, c.valor_presupuestado, c.valor_real,"
                             " (c.valor_real - c.valor_presupuestado) AS diferencia FROM cif c LEFT JOIN centros_costo cc ON cc.id=c.centro_costo_id"
                             " WHERE c.empresa_id=%s AND c.periodo_id=%s ORDER BY c.comportamiento, c.codigo", (eid, periodo_id))
        S.append({"titulo": "CIF presupuestados y reales", "columnas": [
            col("codigo", "Código"), col("concepto", "Concepto"), col("comportamiento", "Comportamiento"),
            col("centro", "Centro de costo"), col("valor_presupuestado", "Presupuestado", "moneda"),
            col("valor_real", "Real", "moneda"), col("diferencia", "Diferencia", "moneda")],
            "filas": filas, "totales": {"codigo": "TOTAL", **_sumar(filas, ["valor_presupuestado", "valor_real", "diferencia"])}})
        v = calculos.variaciones_periodo(eid, periodo_id)["cif"]
        etiquetas = [("base_aplicacion", "Base de aplicación"), ("nivel_normal", "Nivel normal"),
                     ("cif_fijo_presupuestado", "CIF fijo presupuestado"), ("cif_variable_presupuestado", "CIF variable presupuestado"),
                     ("tasa_fija", "Tasa fija"), ("tasa_variable", "Tasa variable"), ("tasa_total", "Tasa predeterminada"),
                     ("base_estandar_permitida", "Base estándar permitida"), ("presupuesto_flexible", "Presupuesto flexible"),
                     ("cif_aplicado", "CIF aplicado"), ("cif_real", "CIF real"), ("variacion_total_cif", "Variación total de CIF")]
        S.append({"titulo": "Tasa predeterminada y aplicación", "columnas": [col("c", "Concepto"), col("v", "Valor", "numero")],
                  "filas": [{"c": t, "v": v[k]} for k, t in etiquetas]})

    elif tipo == "produccion":
        filas = []
        for h in _ordenes(eid, periodo_id, producto_id):
            o = h["orden"]
            filas.append({"numero": o["numero"], "fecha": o["fecha"], "producto": o["producto"], "estado": o["estado"],
                          "terminadas": o["cantidad_terminada"], "en_proceso": o["cantidad_en_proceso"],
                          "costo_estandar": h["total_estandar"], "costo_real": h["total_real"],
                          "costo_terminado": h["costo_terminado_estandar"], "costo_en_proceso": h["costo_en_proceso_estandar"],
                          "diferencia": h["diferencia_total"]})
        S.append({"titulo": "Órdenes de producción", "columnas": [
            col("numero", "Orden"), col("fecha", "Fecha"), col("producto", "Producto"), col("estado", "Estado"),
            col("terminadas", "Terminadas", "numero"), col("en_proceso", "En proceso", "numero"),
            col("costo_estandar", "Costo estándar", "moneda"), col("costo_real", "Costo real", "moneda"),
            col("costo_terminado", "Prod. terminada (est.)", "moneda"), col("costo_en_proceso", "Prod. en proceso (est.)", "moneda"),
            col("diferencia", "Diferencia", "moneda")],
            "filas": filas, "totales": {"numero": "TOTAL", **_sumar(filas, ["terminadas", "en_proceso", "costo_estandar", "costo_real", "costo_terminado", "costo_en_proceso", "diferencia"])}})

    elif tipo == "variaciones":
        v = calculos.variaciones_periodo(eid, periodo_id, producto_id)
        S.append({"titulo": "Resumen de variaciones del periodo", "columnas": [
            col("elemento", "Elemento"), col("tipo", "Variación"), col("valor", "Valor", "moneda"), col("naturaleza", "Naturaleza")],
            "filas": v["variaciones"],
            "totales": {"elemento": "TOTAL", "valor": v["variacion_total"], "naturaleza": v["naturaleza_total"]}})
        S.append({"titulo": "Favorables y desfavorables", "columnas": [col("c", "Concepto"), col("v", "Valor", "moneda")],
                  "filas": [{"c": "Total variaciones favorables", "v": v["total_favorable"]},
                            {"c": "Total variaciones desfavorables", "v": v["total_desfavorable"]},
                            {"c": "Variación neta", "v": v["variacion_total"]}]})
        det = [{"orden": o["numero"], **x} for o in v["ordenes"] for x in o["variaciones"] if abs(x["valor"]) >= 0.005]
        S.append({"titulo": "Detalle por orden", "columnas": [
            col("orden", "Orden"), col("tipo", "Variación"), col("referencia", "Referencia"),
            col("valor", "Valor", "moneda"), col("naturaleza", "Naturaleza")], "filas": det})

    elif tipo == "estado_costos":
        ords = _ordenes(eid, periodo_id, producto_id)
        mp_r = sum(h["resumen"][0]["real_total"] for h in ords)
        mo_r = sum(h["resumen"][1]["real_total"] for h in ords)
        cif_r = sum(h["resumen"][2]["real_total"] for h in ords)
        mp_s = sum(h["resumen"][0]["estandar_total"] for h in ords)
        mo_s = sum(h["resumen"][1]["estandar_total"] for h in ords)
        cif_s = sum(h["resumen"][2]["estandar_total"] for h in ords)
        tot_r, tot_s = mp_r + mo_r + cif_r, mp_s + mo_s + cif_s
        wip_s = sum(h["costo_en_proceso_estandar"] for h in ords)
        wip_r = sum(h["total_real"] * (h["costo_en_proceso_estandar"] / h["total_estandar"]) if h["total_estandar"] else 0 for h in ords)
        r = calculos.r2
        S.append({"titulo": "Estado de costos de producción", "columnas": [
            col("c", "Concepto"), col("real", "Costo real", "moneda"), col("estandar", "Costo estándar", "moneda"),
            col("diferencia", "Diferencia", "moneda")],
            "filas": [
                {"c": "Materia prima directa utilizada", "real": r(mp_r), "estandar": r(mp_s), "diferencia": r(mp_r - mp_s)},
                {"c": "(+) Mano de obra directa", "real": r(mo_r), "estandar": r(mo_s), "diferencia": r(mo_r - mo_s)},
                {"c": "(+) Costos indirectos de fabricación", "real": r(cif_r), "estandar": r(cif_s), "diferencia": r(cif_r - cif_s)},
                {"c": "(=) Costo de producción del periodo", "real": r(tot_r), "estandar": r(tot_s), "diferencia": r(tot_r - tot_s)},
                {"c": "(+) Inventario inicial de productos en proceso", "real": 0, "estandar": 0, "diferencia": 0},
                {"c": "(-) Inventario final de productos en proceso", "real": r(wip_r), "estandar": r(wip_s), "diferencia": r(wip_r - wip_s)},
                {"c": "(=) Costo de la producción terminada", "real": r(tot_r - wip_r), "estandar": r(tot_s - wip_s),
                 "diferencia": r((tot_r - wip_r) - (tot_s - wip_s))},
            ]})
        rep["nota"] = ("El CIF real corresponde a la porción del CIF del periodo prorrateada a las órdenes según la base real. "
                       "El inventario inicial en proceso se considera cero porque cada orden se costea dentro de su periodo.")
    return rep


# ---------------------------------------------------------------------
# Exportación
# ---------------------------------------------------------------------
def _fmt(v, tipo):
    if v is None:
        return ""
    if tipo in ("moneda", "numero") and isinstance(v, (int, float)):
        dec = cfg.entero("sistema", "decimales", 2)
        texto = "{:,.{d}f}".format(v, d=dec if tipo == "moneda" else (2 if float(v) != int(v) else 0))
        return texto.replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


def a_excel(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rep["titulo"][:31]
    azul = PatternFill("solid", fgColor="1F3A5F")
    gris = PatternFill("solid", fgColor="E8EDF3")
    borde = Border(bottom=Side(style="thin", color="C9D3DF"))
    ws.append([rep["empresa"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["NIT %s" % rep["nit"]])
    ws.append([rep["titulo"]])
    ws["A3"].font = Font(bold=True, size=12, color="1F3A5F")
    ws.append([rep.get("subtitulo", "")])
    ws.append(["Generado: %s" % rep["generado"]])
    fila = 7
    anchos = {}
    for s in rep["secciones"]:
        ws.cell(row=fila, column=1, value=s["titulo"]).font = Font(bold=True, size=11, color="1F3A5F")
        fila += 1
        for j, c in enumerate(s["columnas"], 1):
            cel = ws.cell(row=fila, column=j, value=c["titulo"])
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = azul
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            anchos[j] = max(anchos.get(j, 10), len(c["titulo"]) + 2)
        fila += 1
        for f in s["filas"]:
            for j, c in enumerate(s["columnas"], 1):
                v = f.get(c["clave"])
                cel = ws.cell(row=fila, column=j, value=v)
                if c["tipo"] in ("moneda", "numero") and isinstance(v, (int, float)):
                    cel.number_format = '#,##0.00' if c["tipo"] == "moneda" else '#,##0.####'
                cel.border = borde
                anchos[j] = max(anchos.get(j, 10), min(len(str(v or "")) + 2, 45))
            fila += 1
        if s.get("totales"):
            for j, c in enumerate(s["columnas"], 1):
                v = s["totales"].get(c["clave"])
                cel = ws.cell(row=fila, column=j, value=v)
                cel.font = Font(bold=True)
                cel.fill = gris
                if isinstance(v, (int, float)):
                    cel.number_format = '#,##0.00'
            fila += 1
        fila += 1
    if rep.get("nota"):
        ws.cell(row=fila, column=1, value="Nota: " + rep["nota"]).font = Font(italic=True, size=9)
    for j, w in anchos.items():
        ws.column_dimensions[get_column_letter(j)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def a_pdf(rep):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=rep["titulo"])
    est = getSampleStyleSheet()
    azul = colors.HexColor("#1F3A5F")
    celda = ParagraphStyle("celda", parent=est["Normal"], fontSize=7.5, leading=9)
    celda_d = ParagraphStyle("celda_d", parent=celda, alignment=2)
    cab = ParagraphStyle("cab", parent=celda, textColor=colors.white, fontName="Helvetica-Bold", alignment=1)
    hist = [Paragraph("<b>%s</b> · NIT %s" % (rep["empresa"], rep["nit"]), est["Normal"]),
            Paragraph(rep["titulo"], ParagraphStyle("t", parent=est["Title"], textColor=azul, alignment=0, fontSize=16)),
            Paragraph("%s · Generado %s" % (rep.get("subtitulo", ""), rep["generado"]), est["Normal"]), Spacer(1, 8)]
    ancho_total = landscape(A4)[0] - 2.8 * cm
    for s in rep["secciones"]:
        hist.append(Paragraph("<b>%s</b>" % s["titulo"], ParagraphStyle("s", parent=est["Heading4"], textColor=azul)))
        data = [[Paragraph(c["titulo"], cab) for c in s["columnas"]]]
        for f in s["filas"]:
            data.append([Paragraph(_fmt(f.get(c["clave"]), c["tipo"]), celda_d if c["tipo"] in ("moneda", "numero") else celda)
                         for c in s["columnas"]])
        if s.get("totales"):
            data.append([Paragraph("<b>%s</b>" % _fmt(s["totales"].get(c["clave"]), c["tipo"]),
                                   celda_d if c["tipo"] in ("moneda", "numero") else celda) for c in s["columnas"]])
        n = len(s["columnas"])
        primera = 0.26 if n > 3 else 0.5
        anchos = [ancho_total * primera] + [ancho_total * (1 - primera) / (n - 1)] * (n - 1) if n > 1 else [ancho_total]
        t = Table(data, colWidths=anchos, repeatRows=1)
        estilo = [("BACKGROUND", (0, 0), (-1, 0), azul), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                  ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D3DF")),
                  ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")])]
        if s.get("totales"):
            estilo.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8EDF3")))
        t.setStyle(TableStyle(estilo))
        hist += [t, Spacer(1, 10)]
    if rep.get("nota"):
        hist.append(Paragraph("<i>Nota: %s</i>" % rep["nota"], celda))
    doc.build(hist)
    buf.seek(0)
    return buf
