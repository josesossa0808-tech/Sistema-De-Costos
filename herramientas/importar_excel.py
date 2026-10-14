"""Importa al Sistema de Costos la plantilla de costeo en Excel (formato "Costeo Deligaby").

Lee las hojas «Insumos» y «Postre» y crea, a través de la propia API del sistema
(con todas sus validaciones y su registro de actividades):

  - la empresa y un periodo contable
  - unidades de medida, centro de costo y proveedores
  - materias primas (insumos) con su precio estándar = costo unitario del Excel
  - productos con su precio de venta
  - estándar de materiales por producto (la receta por unidad)
  - cargo de mano de obra y estándar de horas por producto
  - CIF del periodo (mantenimiento de equipos y servicios) y su base de aplicación

Funciona con el motor configurado en config/config.ini (SQLite o MySQL).

Uso (desde la carpeta del sistema):
    python herramientas/importar_excel.py                       (usa datos/Costeo_Deligaby.xlsx)
    python herramientas/importar_excel.py ruta/al/archivo.xlsx
    python herramientas/importar_excel.py --reemplazar          (borra la empresa y la vuelve a cargar)
    python herramientas/importar_excel.py --usuario admin --contrasena "MiClave*"
"""
import argparse
import os
import re
import sys
import unicodedata

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "backend"))

import openpyxl  # noqa: E402

import configuracion as cfg  # noqa: E402

# ---------------------------------------------------------------------
# Parámetros de la plantilla (tomados de las notas del Excel)
# ---------------------------------------------------------------------
UNIDADES = {  # unidad del Excel -> (código, nombre, tipo)
    "gramo": ("GR", "Gramo", "Masa"), "kilo": ("KG", "Kilogramo", "Masa"),
    "onza": ("OZ", "Onza", "Masa"), "libra": ("LB", "Libra", "Masa"),
    "mililitro": ("ML", "Mililitro", "Volumen"), "litro": ("LT", "Litro", "Volumen"),
    "galon": ("GAL", "Galón", "Volumen"), "unidad": ("UND", "Unidad", "Conteo"),
    "paquete": ("PAQ", "Paquete", "Conteo"), "rollo": ("ROL", "Rollo", "Conteo"),
    "metro": ("MT", "Metro", "Longitud"),
}
# Proveedores según la nota de la hoja Insumos (fuentes de los PDF de cotización)
PROVEEDORES = {
    "Almacenes Éxito": ["harina", "azucar", "azucar glas", "mantequilla", "huevos", "leche entera", "polvo para hornear"],
    "Mercacentro": ["esencia de vainilla"],
    "Dispropan Caribe": ["vitina (margarina)"],
    "Orquídea": ["capacillo"],
    "Los Cartujos": ["caja x12 cupcakes"],
}
TARIFA_HORA_MOD = 10000  # Nota 3 de la hoja Postre: mano de obra 1,5 h x $10.000


def simple(texto):
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def num(v, defecto=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return defecto


# ---------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------
def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    for hoja in ("Insumos", "Postre"):
        if hoja not in wb.sheetnames:
            raise SystemExit("El archivo no tiene la hoja «%s». ¿Es la plantilla de costeo correcta?" % hoja)

    # --- Insumos
    ws = wb["Insumos"]
    filas = list(ws.iter_rows(values_only=True))
    ini = next((i for i, f in enumerate(filas) if simple(f[1]) == "nombre"), None)
    if ini is None:
        raise SystemExit("No se encontró la tabla de insumos (columna «Nombre») en la hoja Insumos.")
    insumos = []
    for f in filas[ini + 1:]:
        nombre = f[1]
        if not nombre or simple(nombre).startswith(("nota", "precios", "fuentes", "unidades")):
            if insumos:
                break
            continue
        costo_compra, cantidad = num(f[3]), num(f[4])
        unitario = num(f[9]) or num(f[7]) or (costo_compra / cantidad if cantidad else 0)
        insumos.append({"nombre": str(nombre).strip(), "costo_compra": costo_compra, "cantidad_compra": cantidad,
                        "unidad": str(f[5] or "Unidad").strip(), "merma": num(f[6]), "costo_unitario": unitario})

    # --- Postre (productos, recetas y gastos)
    ws = wb["Postre"]
    filas = list(ws.iter_rows(values_only=True))
    ini = next((i for i, f in enumerate(filas) if simple(f[1]) == "nombre de platillo"), None)
    if ini is None:
        raise SystemExit("No se encontró la tabla de productos («Nombre de Platillo») en la hoja Postre.")
    enc = [simple(x) for x in filas[ini]]

    def col(nombre):
        return enc.index(nombre) if nombre in enc else None

    gastos = {k: col(k) for k in ("mano de obra", "servicios", "oper. domicilio", "sueldos admin.", "depreciacion", "otros gastos")}
    productos = []
    for f in filas[ini + 1:]:
        nombre = f[1]
        if not nombre:
            if productos and not any(f[2:]):
                if f[0] is None:
                    break
            continue
        if simple(nombre).startswith(("costos por", "notas")):
            break
        receta = []
        c = 3
        while c + 2 < len(f) and enc[c].startswith("insumo"):
            if f[c] and num(f[c + 2]) > 0:
                receta.append({"insumo": str(f[c]).strip(), "cantidad": num(f[c + 2])})
            c += 3
        g = {k: (num(f[i]) if i is not None else 0) for k, i in gastos.items()}
        es_cupcake = any(p in simple(nombre) for p in ("cupcake", "copcake", "cup cake"))
        productos.append({"nombre": str(nombre).strip(), "precio": num(f[2]), "receta": receta, "gastos": g,
                          "unidades_lote": 30 if es_cupcake else 1})

    # --- Bloques de mantenimiento y reparto
    mantenimiento, producciones = [], None
    for f in filas:
        for i, v in enumerate(f):
            s = simple(v)
            if s in ("horno", "batidora") and i + 1 < len(f) and num(f[i + 1]) > 0:
                mantenimiento.append((str(v).strip(), num(f[i + 1])))
            if s == "producciones que cubre" and i + 1 < len(f):
                producciones = num(f[i + 1]) or None
    return insumos, productos, mantenimiento, producciones or 3


# ---------------------------------------------------------------------
# Carga por la API del sistema
# ---------------------------------------------------------------------
class Api:
    def __init__(self, cliente, token):
        self.c, self.h = cliente, {"Authorization": "Bearer " + token}

    def empresa(self, eid):
        self.h["X-Empresa"] = str(eid)

    def _r(self, r, que):
        if r.status_code >= 400:
            raise SystemExit("Error al crear %s: %s" % (que, (r.get_json() or {}).get("error") or r.data[:300]))
        return r.get_json()

    def crear(self, recurso, datos, que=None):
        return self._r(self.c.post("/api/r/" + recurso, json=datos, headers=self.h), que or recurso)["id"]

    def listar(self, recurso):
        return self._r(self.c.get("/api/r/" + recurso, headers=self.h), recurso)

    def eliminar(self, recurso, rid):
        return self._r(self.c.delete("/api/r/%s/%s" % (recurso, rid), headers=self.h), recurso)

    def get(self, ruta):
        return self._r(self.c.get(ruta, headers=self.h), ruta)


def importar(args):
    insumos, productos, mantenimiento, producciones = leer_excel(args.archivo)
    print("Excel leído: %d insumos, %d productos." % (len(insumos), len(productos)))

    import app as aplicacion  # se importa aquí para respetar SC_CONFIG
    flask_app = aplicacion.inicializar_bd()
    cli = flask_app.test_client()
    usuario = args.usuario or cfg.obtener("administrador_inicial", "usuario", "admin")
    clave = args.contrasena or cfg.obtener("administrador_inicial", "contrasena", "Admin123*")
    r = cli.post("/api/auth/login", json={"usuario": usuario, "contrasena": clave})
    if r.status_code != 200:
        raise SystemExit("No se pudo iniciar sesión como «%s»: %s\nUse --usuario y --contrasena con un administrador."
                         % (usuario, (r.get_json() or {}).get("error")))
    api = Api(cli, r.get_json()["token"])

    # --- Empresa (si ya existe, se reemplaza solo con --reemplazar)
    existente = [e for e in api.listar("empresas") if simple(e["razon_social"]) == simple(args.empresa)]
    if existente:
        if not args.reemplazar:
            raise SystemExit("La empresa «%s» ya existe en el sistema. Ejecute con --reemplazar para borrarla y "
                             "volver a cargar el Excel." % args.empresa)
        for e in existente:
            api.eliminar("empresas", e["id"])
        print("Empresa anterior eliminada.")
    eid = api.crear("empresas", {"nit": args.nit, "razon_social": args.empresa, "nombre_comercial": args.empresa,
                                 "tipo_persona": "Natural", "moneda": "COP", "actividad_economica": "1081"}, "la empresa")
    api.empresa(eid)

    # --- Periodo
    pid = api.crear("periodos", {"nombre": args.periodo, "fecha_inicio": args.inicio, "fecha_fin": args.fin}, "el periodo")

    # --- Unidades de medida (catálogo general: se reutilizan si ya existen)
    unidades = {simple(u["codigo"]): u["id"] for u in api.listar("unidades_medida")}

    def unidad_id(nombre_excel):
        cod, nom, tipo = UNIDADES.get(simple(nombre_excel), (simple(nombre_excel)[:10].upper(), str(nombre_excel), "Conteo"))
        if simple(cod) not in unidades:
            unidades[simple(cod)] = api.crear("unidades_medida", {"codigo": cod, "nombre": nom, "tipo": tipo}, "la unidad " + cod)
        return unidades[simple(cod)]

    und = unidad_id("Unidad")

    # --- Centro de costo y proveedores
    centro = api.crear("centros_costo", {"codigo": "PROD", "nombre": "Cocina de producción", "tipo": "Producción"})
    prov_de = {}
    for i, (nombre, items) in enumerate(PROVEEDORES.items(), 1):
        if any(simple(x["nombre"]) in items for x in insumos):
            prov = api.crear("proveedores", {"nit": "SIN-NIT-%02d" % i, "nombre": nombre}, "el proveedor " + nombre)
            for it in items:
                prov_de[it] = prov

    # --- Materias primas
    mp = {}
    for i, x in enumerate(insumos, 1):
        mp[simple(x["nombre"])] = api.crear("materias_primas", {
            "codigo": "MP%02d" % i, "descripcion": x["nombre"], "unidad_id": unidad_id(x["unidad"]),
            "precio_estandar": round(x["costo_unitario"], 4), "proveedor_id": prov_de.get(simple(x["nombre"])),
            "centro_costo_id": centro}, "el insumo " + x["nombre"])
    # «Otros gastos» del Excel = manga pastelera desechable ($500 por lote de cupcakes decorados)
    otros = [p for p in productos if p["gastos"]["otros gastos"] > 0]
    if otros:
        precio_manga = round(otros[0]["gastos"]["otros gastos"] * otros[0]["unidades_lote"], 2)
        mp["manga"] = api.crear("materias_primas", {
            "codigo": "MP%02d" % (len(insumos) + 1), "descripcion": "Manga pastelera desechable",
            "unidad_id": und, "precio_estandar": precio_manga, "centro_costo_id": centro}, "la manga pastelera")

    # --- Mano de obra
    cargo = api.crear("cargos", {"nombre": "Repostero(a)", "tarifa_estandar_hora": TARIFA_HORA_MOD,
                                 "descripcion": "Tarifa según la plantilla de costeo ($10.000 por hora)"})

    # --- Productos, recetas y estándar de mano de obra
    for i, p in enumerate(productos, 1):
        g, lote = p["gastos"], p["unidades_lote"]
        desc = ("Importado del Excel. %s. Domicilio / gastos de operación por unidad: $%s (no es costo de producción)."
                % ("Lote de %d unidades" % lote if lote > 1 else "Se vende por unidad completa",
                   format(round(g["oper. domicilio"], 2), ",.2f")))
        prod = api.crear("productos", {"codigo": "P%02d" % i, "nombre": p["nombre"], "unidad_id": und,
                                       "precio_venta": p["precio"], "horas_maquina_estandar": round(1 / lote, 6),
                                       "descripcion": desc[:300]}, "el producto " + p["nombre"])
        for linea in p["receta"]:
            clave = simple(linea["insumo"])
            if clave not in mp:
                raise SystemExit("El producto «%s» usa «%s», que no está en la hoja Insumos." % (p["nombre"], linea["insumo"]))
            api.crear("estandar_materiales", {"producto_id": prod, "materia_prima_id": mp[clave],
                                              "cantidad_estandar": round(linea["cantidad"], 6)},
                      "la receta de " + p["nombre"])
        if g["otros gastos"] > 0 and "manga" in mp:
            api.crear("estandar_materiales", {"producto_id": prod, "materia_prima_id": mp["manga"],
                                              "cantidad_estandar": round(1 / lote, 6)}, "la manga de " + p["nombre"])
        if g["mano de obra"] > 0:
            api.crear("estandar_mano_obra", {"producto_id": prod, "cargo_id": cargo,
                                             "horas_estandar": round(g["mano de obra"] / TARIFA_HORA_MOD, 6),
                                             "tarifa_estandar": TARIFA_HORA_MOD}, "la mano de obra de " + p["nombre"])

    # --- CIF del periodo: base = horneadas (1 por torta, 1 por lote de 30 cupcakes)
    api.crear("cif_parametros", {"periodo_id": pid, "base_aplicacion": "horas_maquina", "nivel_normal": producciones})
    n = 0
    for equipo, valor in mantenimiento:
        n += 1
        api.crear("cif", {"periodo_id": pid, "codigo": "CIF%02d" % n, "concepto": "Mantenimiento de %s" % equipo.lower(),
                          "comportamiento": "Fijo", "centro_costo_id": centro, "valor_presupuestado": valor})
    tortas = [p for p in productos if p["unidades_lote"] == 1 and p["gastos"]["servicios"] > 0]
    servicios = tortas[0]["gastos"]["servicios"] if tortas else 0
    if servicios:
        n += 1
        api.crear("cif", {"periodo_id": pid, "codigo": "CIF%02d" % n, "concepto": "Servicios (agua, luz, gas)",
                          "comportamiento": "Variable", "centro_costo_id": centro,
                          "valor_presupuestado": round(servicios * producciones, 2)})

    # --- Producción de referencia (opcional): las 3 producciones que cubre el mantenimiento en el Excel.
    #     El Excel no trae datos reales, así que lo real = lo estándar de la plantilla.
    if not args.sin_produccion:
        trab = api.crear("trabajadores", {
            "documento": "TRAB-01", "nombre": "Repostero(a) de %s" % args.empresa, "cargo_id": cargo,
            "centro_costo_id": centro, "salario_mensual": TARIFA_HORA_MOD * 192, "horas_mes": 192}, "el trabajador")
        for x in api.listar("cif"):
            datos = {k: x[k] for k in ("periodo_id", "codigo", "concepto", "comportamiento", "centro_costo_id", "valor_presupuestado")}
            datos["valor_real"] = x["valor_presupuestado"]
            r = api.c.put("/api/r/cif/%s" % x["id"], json=datos, headers=api.h)
            api._r(r, "el CIF real")
        por_nombre = {simple(p["nombre"]): p for p in api.listar("productos")}
        elegidos = [p for p in productos if p["unidades_lote"] == 1][:2] + [p for p in productos if p["unidades_lote"] > 1][:1]
        for i, p in enumerate(elegidos[:int(producciones)], 1):
            prod = por_nombre[simple(p["nombre"])]
            est = api.get("/api/estandar/%s?periodo_id=%s" % (prod["id"], pid))
            q = p["unidades_lote"]
            r = api.c.post("/api/ordenes", json={
                "periodo_id": pid, "producto_id": prod["id"], "numero": "OP-%04d" % i, "centro_costo_id": centro,
                "fecha": "%s-%02d" % (args.inicio[:7], 5 * i), "cantidad_terminada": q, "horas_maquina_reales": 1,
                "observaciones": "Producción de referencia tomada del Excel (consumo real = receta estándar)",
                "materiales": [{"materia_prima_id": m["materia_prima_id"], "cantidad_real": round(m["cantidad_estandar"] * q, 6),
                                "precio_real": m["precio_estandar"]} for m in est["materiales"]],
                "mano_obra": [{"trabajador_id": trab, "horas_reales": round(sum(x["horas_estandar"] for x in est["mano_obra"]) * q, 4),
                               "tarifa_real": TARIFA_HORA_MOD}],
            }, headers=api.h)
            api._r(r, "la orden de " + p["nombre"])
        print("Producción de referencia: %d órdenes registradas." % len(elegidos[:int(producciones)]))

    # --- Resumen
    print("\nImportación terminada: empresa «%s», periodo «%s»." % (args.empresa, args.periodo))
    print("%-42s %14s %14s %14s %14s" % ("Producto", "Materiales", "Mano de obra", "CIF", "Costo estándar"))
    for p in api.listar("productos"):
        e = api.get("/api/estandar/%s?periodo_id=%s" % (p["id"], pid))
        print("%-42s %14s %14s %14s %14s" % (p["nombre"][:42], *(format(e[k], ",.2f") for k in
                                              ("mp_unitario", "mod_unitario", "cif_unitario", "total_unitario"))))
    return api, pid


def main():
    ap = argparse.ArgumentParser(description="Importa la plantilla de costeo en Excel al Sistema de Costos")
    ap.add_argument("archivo", nargs="?", default=os.path.join(RAIZ, "datos", "Costeo_Deligaby.xlsx"))
    ap.add_argument("--empresa", default="Deligaby")
    ap.add_argument("--nit", default="PENDIENTE", help="NIT de la empresa (se puede corregir luego en el sistema)")
    ap.add_argument("--periodo", default="Octubre 2026")
    ap.add_argument("--inicio", default="2026-10-01")
    ap.add_argument("--fin", default="2026-10-31")
    ap.add_argument("--usuario")
    ap.add_argument("--contrasena")
    ap.add_argument("--reemplazar", action="store_true", help="Borra la empresa si ya existe y la vuelve a cargar")
    ap.add_argument("--sin-produccion", action="store_true",
                    help="No registra el trabajador, el CIF real ni las órdenes de producción de referencia")
    args = ap.parse_args()
    if not os.path.exists(args.archivo):
        raise SystemExit("No se encontró el archivo: %s" % args.archivo)
    importar(args)


if __name__ == "__main__":
    main()
